"""Real puzzle UI -> durable outbox -> isolated Firebase -> rendered leaderboard.

Run inside firebase emulators:exec --project demo-rudn --only auth,database.
Only config and a read-only engine geometry observer are injected. The real
backend, authentication, completion bridge, rules and leaderboard remain active.
"""
from __future__ import annotations
import argparse
import asyncio
import hashlib
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
import json
import os
from pathlib import Path
import re
import sys
import threading
import time
import urllib.request
import openpyxl
from urllib.parse import urlsplit
from playwright.async_api import async_playwright
from test_puzzle_catalog import ROOT, READ_ONLY_HOOK, ready, place_pieces, trusted_drop

NAMESPACE='demo-rudn-default-rtdb'
DATABASE='http://127.0.0.1:9000'
BACKEND_SPEC=re.search(r"from './(backend\.js[^']*)'",(ROOT/'site/assets/js/puzzle-bootstrap.js').read_text('utf8')).group(1)

def db(path='',value=None,method='GET'):
    url=DATABASE+'/'+path+'.json?ns='+NAMESPACE
    assert urlsplit(url).hostname=='127.0.0.1'
    request=urllib.request.Request(url,method=method,headers={'Content-Type':'application/json','Authorization':'Bearer owner'},data=json.dumps(value).encode() if method!='GET' else None)
    with urllib.request.build_opener(urllib.request.ProxyHandler({})).open(request,timeout=10) as response:return json.load(response)

def create_teacher():
    request=urllib.request.Request('http://127.0.0.1:9099/identitytoolkit.googleapis.com/v1/accounts:signUp?key=demo-rudn',method='POST',headers={'Content-Type':'application/json'},data=json.dumps({'email':'omnistat@yandex.ru','password':'Emulator-Only-Teacher-2026','returnSecureToken':True}).encode())
    try:
        with urllib.request.build_opener(urllib.request.ProxyHandler({})).open(request,timeout=10) as response:json.load(response)
    except urllib.error.HTTPError as error:
        body=json.load(error)
        assert body.get('error',{}).get('message')=='EMAIL_EXISTS',body

class EmulatorSite:
    def __enter__(self):
        self.offline=False
        server=self
        class Handler(SimpleHTTPRequestHandler):
            def __init__(self,*a,**kw):super().__init__(*a,directory=str(ROOT/'site'),**kw)
            def log_message(self,*_):pass
            def end_headers(self):
                self.send_header('Content-Security-Policy',"connect-src 'self' http://127.0.0.1:9000 http://127.0.0.1:9099 ws://127.0.0.1:9000; frame-src 'none'")
                super().end_headers()
            def copyfile(self,source,output):
                try:super().copyfile(source,output)
                except (BrokenPipeError,ConnectionResetError,ConnectionAbortedError):pass
            def do_GET(self):
                if server.offline:
                    self.send_response(503);self.send_header('Cache-Control','no-store');self.end_headers();self.wfile.write(b'Isolated site transport outage');return
                name=urlsplit(self.path).path
                if not name.startswith('/RUDN/'):
                    self.send_error(404);return
                self.path=self.path[len('/RUDN'):]
                if name.endswith('/assets/js/config.js') or name.endswith('/assets/js/puzzle-engine.js'):
                    source=(ROOT/'site'/name[len('/RUDN/'):]).read_text('utf8')
                    if name.endswith('/config.js'):
                        source+="\nCONFIG.emulators={host:'127.0.0.1',databasePort:9000,auth:'http://127.0.0.1:9099'};\n"
                    else:
                        marker='  function checkpoint() {'
                        assert source.count(marker)==1
                        source=source.replace(marker,READ_ONLY_HOOK+'\n'+marker,1)
                    body=source.encode('utf8');self.send_response(200)
                    self.send_header('Content-Type','text/javascript; charset=utf-8');self.send_header('Content-Length',str(len(body)));self.end_headers();self.wfile.write(body)
                else:super().do_GET()
        self.httpd=ThreadingHTTPServer(('127.0.0.1',0),Handler)
        self.thread=threading.Thread(target=self.httpd.serve_forever,daemon=True);self.thread.start()
        self.base=f'http://127.0.0.1:{self.httpd.server_port}/RUDN/'
        return self
    def __exit__(self,*_):self.httpd.shutdown();self.httpd.server_close();self.thread.join(timeout=2)

async def backend(page,expression):
    return await page.evaluate("async ({base,spec,expression})=>{const {backend}=await import(base+'assets/js/'+spec);return await (new Function('backend','return ('+expression+')'))(backend)}",{'base':page._qa_base,'spec':BACKEND_SPEC,'expression':expression})

async def wait_public(attempt_id):
    for _ in range(100):
        row=await asyncio.to_thread(db,'results/'+attempt_id)
        if row:return row
        await asyncio.sleep(.2)
    raise AssertionError('Completed UI result did not reach emulator: '+attempt_id)

async def free_module_contracts(page,ticket):
    """Real shared backend/rules, paired with the separate full module UI suite."""
    before={slug:{'points':4,'max':5,'updatedAt':1,'sourceAttemptId':'isolated-existing-grade'} for slug in ('seminar-5','seminar-7')}
    for slug,grade in before.items():db(f'rudn-platform/v1/grades/{ticket}/{slug}',grade,'PUT')
    ids=await backend(page,"""(async()=>{
      const suffix=crypto.randomUUID(),ids={};
      for(const [slug,type] of [['governor-freeplay','governor-simulator'],['career-freeplay','career-diagnostic']]){
        const id='free-contract-'+slug+'-'+suffix;ids[slug]=id;
        await backend.saveAttempt({id,activitySlug:slug,type,recordGrade:false,
          draftMode:slug==='governor-freeplay'?'campaign':'diagnostic',points:4,maxPoints:5,createdAt:new Date().toISOString()});
      }
      ids['reception-freeplay']='free-contract-reception-'+suffix;
      await backend.checkpoint({owner:'student:'+backend.getProfile().studentKey,
        activitySlug:'reception-freeplay',mode:'2026-2027:practice',attemptId:ids['reception-freeplay'],
        state:{phase:'completed',practice:true,score:5}});
      await backend.syncLocalToCloud();return ids;
    })()""")
    for slug,id in ids.items():
        path=f'rudn-platform/v1/checkpoints/{ticket}/{slug}/{id}' if slug=='reception-freeplay' else f'rudn-platform/v1/attempts/{ticket}/{id}'
        for _ in range(100):
            value=await asyncio.to_thread(db,path)
            if value:break
            await asyncio.sleep(.2)
        assert value,('Free record did not sync',slug)
        if slug=='reception-freeplay':
            assert value['mode']=='2026-2027:practice' and value['current']['state']['phase']=='completed'
            assert db(f'rudn-platform/v1/attempts/{ticket}/{id}') is None
        else:assert value['recordGrade'] is False and value['activitySlug']==slug
        assert db(f'rudn-platform/v1/grades/{ticket}/{slug}') is None
    for slug,grade in before.items():assert db(f'rudn-platform/v1/grades/{ticket}/{slug}')==grade
    assert db(f'rudn-platform/v1/grades/{ticket}/seminar-2')['points']==5
    return {'slugs':list(ids),'synced':True,'freeGradesCreated':0,'priorCourseGrades':{'seminar-2':5,'seminar-5':4,'seminar-7':4}}

async def select_russia(page,server,mode='free'):
    await page.goto(server.base+'apps/puzzle.html?context='+mode,wait_until='domcontentloaded')
    await ready(page)
    if mode=='free':
        await page.locator('[data-puzzle-mode="country-regions"]').click()
        await page.locator('#puzzleCountry').select_option('RUS')
        await page.wait_for_function("window.__puzzleRead().ready&&!window.__puzzleRead().loading&&window.__puzzleRead().mode==='russia-subjects'")
    if (await page.evaluate('window.__puzzleRead()'))['finished']:
        await page.locator('#puzzlePlayAgain').click()
        await page.wait_for_function('!window.__puzzleRead().finished&&!window.__puzzleRead().loading')
    assert (await page.evaluate('window.__puzzleRead()'))['total']==89

async def finish_map(page):
    await page.locator('#puzzleCanvas').scroll_into_view_if_needed()
    await page.wait_for_function("()=>{const r=document.querySelector('#puzzleCanvas').getBoundingClientRect();return r.top>=0&&r.bottom<=innerHeight}")
    await trusted_drop(page)
    outcome=await place_pieces(page,88)
    assert outcome['actual']==88
    await page.locator('#puzzleResultDialog[open]').wait_for()
    return outcome['state']['attemptId']

async def run(browser_name,output):
    ticket=str(int(time.time()*1000))
    profile={'ticket':ticket,'studentKey':ticket,'email':ticket+'@rudn.ru','fullName':'UI Emulator Student '+browser_name,'group':'ГГУбд-01-26'}
    # Isolated local fixture setup only. Never write to a configured production URL.
    # Each engine starts before the synthetic pagination fixtures are installed.
    # Retaining another engine's 205 faster rows would correctly put the new
    # completed student on page three, outside this immediate-row assertion.
    prior_fixtures={key:None for key in (db('results') or {}) if key.startswith('ui-fixture-')}
    if prior_fixtures:db('results',prior_fixtures,'PATCH')
    db('rudn-platform/v1/access/overrides',{str(year):{'topic-2':{'state':'open','updatedAt':1,'teacherUid':'emulator'}} for year in range(2025,2040)},'PUT')
    record={'browser':browser_name,'profileTicketSuffix':ticket[-4:],'results':[],'pageErrors':[],'unexpectedRequests':[]}
    with EmulatorSite() as server:
        async with async_playwright() as playwright:
            kwargs={'firefox_user_prefs':{'network.proxy.type':0}} if browser_name=='firefox' else {}
            browser=await getattr(playwright,browser_name).launch(headless=True,**kwargs)
            # Keep the Firebase fault injection observable. WebKit dispatches
            # controlled fetches through its worker even when the worker lets
            # them pass, bypassing Playwright routes. Worker behavior is tested
            # separately below with a real server-side asset outage.
            context=await browser.new_context(viewport={'width':1366,'height':1000},locale='ru-RU',service_workers='block',accept_downloads=True)
            await context.add_init_script("""(()=>{if(location.hostname==='127.0.0.1'&&!localStorage.getItem('qa-profile')){localStorage.setItem('rudn.profile.v1',JSON.stringify(%s));localStorage.setItem('rudn.locale','ru');localStorage.setItem('qa-profile','1')}})()"""%json.dumps(profile))
            page=await context.new_page();page._qa_base=server.base
            page.on('pageerror',lambda e:record['pageErrors'].append(str(e)))
            page.on('request',lambda request:record['unexpectedRequests'].append(request.url) if urlsplit(request.url).hostname not in ('127.0.0.1',None) else None)
            page.set_default_timeout(20000)
            try:
                await page.goto(server.base+'apps/puzzle.html?context=seminar',wait_until='domcontentloaded')
                await ready(page)
                await backend(page,'backend.ensureStudentCloud()')
                await page.wait_for_function("async spec=>{const {backend}=await import('../assets/js/'+spec);return Boolean(backend.user?.uid&&backend.user.uid===backend.authClient.currentUser?.uid)}",arg=BACKEND_SPEC)
                record['startup']=await backend(page,"({profile:Boolean(backend.profile),admin:backend.isAdmin(),uid:backend.user?.uid||null,currentUid:backend.authClient.currentUser?.uid||null,authReady:backend.authReady,mode:backend.mode,errorCode:backend.error?.code||null})")
                await backend(page,'backend.ensureCloudProfile()')
                assert await backend(page,'Boolean(backend.getProfile()&&!backend.isAdmin()&&backend.user?.uid)')
                await page.wait_for_function("document.querySelector('#geoPuzzleApp').puzzleProgress.canWrite()")
                for index,(difficulty,points) in enumerate([('hard',5),('medium',4),('easy',3)]):
                    print(f'{browser_name}: graded {difficulty}',flush=True)
                    if index:
                        await page.locator('#puzzlePlayAgain').click();await page.wait_for_function('!window.__puzzleRead().finished&&!window.__puzzleRead().loading')
                    await page.locator(f'[data-puzzle-difficulty="{difficulty}"]').click()
                    await page.wait_for_function('d=>window.__puzzleRead().difficulty===d&&!window.__puzzleRead().loading',arg=difficulty)
                    assert await page.locator('#puzzleHint').is_visible()
                    await page.locator('#puzzleHint').click()
                    assert (await page.evaluate('window.__puzzleRead()'))['hints']==1
                    await trusted_drop(page)
                    if index==0:
                        # Firebase transport fails while the actual app stays usable.
                        await context.route(re.compile(r'^http://127\.0\.0\.1:(9000|9099)/'),lambda route:route.abort('internetdisconnected'))
                        await backend(page,'Promise.resolve(backend.db.goOffline(backend.database))')
                        partial=await place_pieces(page,87)
                        assert partial['actual']==87
                        # An already-pending remote refresh must not delay the
                        # newly completed local row until its five-second timeout.
                        held=asyncio.Event();release=asyncio.Event()
                        async def hold_results(route):
                            held.set();await release.wait()
                            await route.abort('internetdisconnected')
                        held_pattern=re.compile(r'^http://127\.0\.0\.1:9000/results\.json\?')
                        await context.route(held_pattern,hold_results)
                        await page.evaluate("window.dispatchEvent(new Event('online'))")
                        await asyncio.wait_for(held.wait(),10)
                        try:
                            outcome=await place_pieces(page,1)
                            assert outcome['actual']==1
                            await page.locator(f'#leader_{difficulty} tr[data-attempt-id="{outcome["state"]["attemptId"]}"]').wait_for(timeout=4000)
                            record['localRowDuringPendingRefresh']=True
                        finally:
                            release.set()
                            await context.unroute(held_pattern)
                    else:
                        outcome=await place_pieces(page,88)
                        assert outcome['actual']==88
                    await page.locator('#puzzleResultDialog[open]').wait_for()
                    attempt_id=outcome['state']['attemptId']
                    await page.locator(f'#leader_{difficulty} tr[data-attempt-id="{attempt_id}"]').wait_for()
                    if index==0:
                        assert db('results/'+attempt_id) is None
                        await page.reload(wait_until='domcontentloaded');await ready(page)
                        await page.locator(f'#leader_{difficulty} tr[data-attempt-id="{attempt_id}"]').wait_for()
                        await context.unroute(re.compile(r'^http://127\.0\.0\.1:(9000|9099)/'))
                        await backend(page,"(async()=>{backend.db.goOnline(backend.database);const {durableStore}=await import(new URL('../assets/js/durable-store.js',location.href));for(const op of await durableStore.listPending({owner:'student:'+backend.getProfile().studentKey,includeDeferred:true}))await durableStore.retry(op.id,{revision:op.revision});await backend.syncLocalToCloud();window.dispatchEvent(new Event('online'))})()")
                    cloud=await wait_public(attempt_id)
                    assert cloud['participant_id']==hashlib.sha256(('rudn-puzzle-participant-v1:'+ticket).encode()).hexdigest()
                    assert cloud['elapsed_ms']>=cloud['time_ms']>1000
                    private=db(f'rudn-platform/v1/attempts/{ticket}/{attempt_id}')
                    assert private['points']==points and private['hints']==1
                    assert db(f'rudn-platform/v1/grades/{ticket}/seminar-2')['points']==5
                    record['results'].append({'id':attempt_id,'difficulty':difficulty,'points':points,'inputPieces':89,'hints':1,'publicRow':True})
                assert len({x['id'] for x in record['results']})==3
                record['freeModuleContracts']=await free_module_contracts(page,ticket)
                await page.reload(wait_until='domcontentloaded');await ready(page)
                await page.keyboard.press('Escape')
                for result in record['results']:await page.locator(f'#leader_{result["difficulty"]} tr[data-attempt-id="{result["id"]}"]').wait_for()
                assert [await el.get_attribute('id') for el in await page.locator('.puzzle-leaderboard table').all()]==['leader_hard','leader_medium','leader_easy']
                await select_russia(page,server,'free')
                print(f'{browser_name}: signed-in free play',flush=True)
                await page.locator('[data-puzzle-difficulty="hard"]').click()
                await page.wait_for_function("window.__puzzleRead().difficulty==='hard'&&!window.__puzzleRead().loading")
                free_id=await finish_map(page)
                await wait_public(free_id)
                free_attempt=db(f'rudn-platform/v1/attempts/{ticket}/{free_id}')
                assert free_attempt['recordGrade'] is False and free_attempt['activitySlug']=='maps-freeplay'
                assert db(f'rudn-platform/v1/grades/{ticket}/maps-freeplay') is None
                assert db(f'rudn-platform/v1/grades/{ticket}/seminar-2')['points']==5
                record['freePlay']={'attempt':free_id,'publicRow':True,'gradeCreated':False,'inputPieces':89}
                await page.keyboard.press('Escape')
                fixture_group={'chromium':'ГГУбд-06-25','firefox':'ГГУбд-06-24','webkit':'ГГУбд-06-23'}[browser_name]
                fixture_prefix='ui-fixture-'+browser_name+'-'+ticket
                previous={key:None for key in (db('results') or {}) if key.startswith('ui-fixture-'+browser_name+'-')}
                if previous:db('results',previous,'PATCH')
                fixtures={}
                for index in range(205):
                    elapsed=7200123 if index==204 else 10000+index
                    fixture={'fio':'=1+2' if index==204 else 'Student 学生 '+str(index),'group':fixture_group,'difficulty':'hard',
                        'time_ms':min(elapsed,3599000),'elapsed_ms':elapsed,'placed':89,'total':89,'timestamp':1767225600000+index,'user_agent':'isolated-fixture',
                        'participant_id':hashlib.sha256((fixture_prefix+str(index)).encode()).hexdigest()}
                    fixtures[fixture_prefix+'-'+str(index)]=fixture
                db('results',fixtures,'PATCH')
                await page.evaluate("window.dispatchEvent(new Event('online'))")
                await page.locator('#puzzleGroupFilter').click()
                checkbox=page.locator(f'#puzzleGroupFilterOptions input[value="{fixture_group}"]')
                await checkbox.wait_for();await checkbox.check()
                assert await page.locator('#leader_hard tbody tr[data-attempt-id]').count()==100
                await page.locator('[data-leaderboard-pagination="hard"] [data-page-step="1"]').click()
                assert await page.locator('#leader_hard tbody tr:first-child td:first-child').inner_text()=='101'
                focus_kept=await page.evaluate("""()=>new Promise((resolve,reject)=>{
                  const nav=document.querySelector('[data-leaderboard-pagination="hard"]');
                  const timer=setTimeout(()=>{observer.disconnect();reject(new Error('No leaderboard refresh observed'))},6000);
                  const observer=new MutationObserver(()=>{clearTimeout(timer);observer.disconnect();resolve(document.activeElement===nav.querySelector('[data-page-step="1"]'))});
                  observer.observe(nav,{childList:true});window.dispatchEvent(new Event('online'));
                })""")
                assert focus_kept,'Background refresh lost pagination keyboard focus'
                record['paginationFocusOnRefresh']=True
                if await page.locator('#puzzleGroupFilter').get_attribute('open') is None:await page.locator('#puzzleGroupFilter summary').click()
                await checkbox.focus()
                extra_key=fixture_prefix+'-focus-group'
                db('results/'+extra_key,{**next(iter(fixtures.values())),'group':'Extra focus group '+browser_name,'participant_id':hashlib.sha256(extra_key.encode()).hexdigest()},'PUT')
                group_focus=await page.evaluate("""value=>new Promise((resolve,reject)=>{
                  const options=document.querySelector('#puzzleGroupFilterOptions');
                  const timer=setTimeout(()=>{observer.disconnect();reject(new Error('No updated group options observed'))},6000);
                  const observer=new MutationObserver(()=>{clearTimeout(timer);observer.disconnect();resolve(document.activeElement.matches('#puzzleGroupFilterOptions input')&&document.activeElement.value===value)});
                  observer.observe(options,{childList:true});window.dispatchEvent(new Event('online'));
                })""",fixture_group)
                assert group_focus,'New group list lost the existing checkbox keyboard focus'
                db('results',{extra_key:None},'PATCH')
                record['groupFocusOnRefresh']=True
                await page.locator('[data-leaderboard-pagination="hard"] [data-page-step="1"]').click()
                assert await page.locator('#leader_hard tbody tr[data-attempt-id]').count()==5
                assert await page.locator('#leader_hard tbody tr:last-child td:nth-child(4)').inner_text()=='120:00'
                # Export is generated by the actual app button, without replacing XLSX.
                async with page.expect_download() as downloaded:await page.locator('#puzzleLeaderboardExport').click()
                download=await downloaded.value;await download.save_as(output/f'ui-{browser_name}.xlsx')
                assert download.suggested_filename=='rudn-map-leaderboard.xlsx'
                book=openpyxl.load_workbook(output/f'ui-{browser_name}.xlsx')
                assert book.sheetnames==['Высокая','Средняя','Низкая']
                assert book.worksheets[0].max_row==206 and book.worksheets[1].max_row==book.worksheets[2].max_row==1
                assert book.worksheets[0]['B206'].value=='=1+2' and book.worksheets[0]['B206'].data_type=='s'
                assert abs(book.worksheets[0]['D206'].value.total_seconds()*1000-7200123)<1
                book.close()
                record['filteredExport']={'groupYear':fixture_group[-2:],'rows':205,'visibleLastPageRows':5,'sheets':3,'literalFormulaText':True,'longDurationPreserved':True}
                print(f'{browser_name}: filtered XLSX passed; offline cache',flush=True)
                await context.route(re.compile(r'^http://127\.0\.0\.1:(9000|9099)/'),lambda route:route.abort('internetdisconnected'))
                await page.reload(wait_until='domcontentloaded');await ready(page)
                await page.keyboard.press('Escape')
                await page.locator('#puzzleGroupFilter').click();await page.locator(f'#puzzleGroupFilterOptions input[value="{fixture_group}"]').check()
                assert await page.locator('#leader_hard tbody tr[data-attempt-id]').count()==100
                async with page.expect_download() as offline_download:await page.locator('#puzzleLeaderboardExport').click()
                await (await offline_download.value).save_as(output/f'ui-{browser_name}-offline.xlsx')
                offline_book=openpyxl.load_workbook(output/f'ui-{browser_name}-offline.xlsx')
                assert offline_book.sheetnames==['Высокая','Средняя','Низкая']
                assert offline_book.worksheets[0].max_row==206 and offline_book.worksheets[1].max_row==offline_book.worksheets[2].max_row==1
                assert offline_book.worksheets[0]['B206'].value=='=1+2' and offline_book.worksheets[0]['B206'].data_type=='s'
                offline_book.close()
                record['cachedRowsAndOfflineExport']=True
                await context.unroute(re.compile(r'^http://127\.0\.0\.1:(9000|9099)/'))
                await page.locator('#puzzleGroupFilterClear').click()
                await page.locator('#puzzleGroupFilter summary').click()
                await page.locator('.puzzle-leaderboard').screenshot(path=str(output/f'leaderboard-{browser_name}.png'))
                await context.close()
                original_result_ids=set(db('results') or {})
                record['excludedRoles']=[]
                for role in ('guest','teacher'):
                    print(f'{browser_name}: {role} result exclusion',flush=True)
                    preview=await browser.new_context(viewport={'width':1366,'height':1000},locale='ru-RU',service_workers='block')
                    preview_page=await preview.new_page();preview_page._qa_base=server.base
                    preview_page.on('pageerror',lambda error:record['pageErrors'].append(str(error)))
                    preview_page.on('request',lambda request:record['unexpectedRequests'].append(request.url) if urlsplit(request.url).hostname not in ('127.0.0.1',None) else None)
                    try:
                        await preview_page.goto(server.base+'apps/puzzle.html?context=free',wait_until='domcontentloaded');await ready(preview_page)
                        if role=='teacher':
                            create_teacher()
                            await backend(preview_page,"backend.adminSignIn('omnistat@yandex.ru','Emulator-Only-Teacher-2026').then(()=>true)")
                            await select_russia(preview_page,server,'seminar')
                            assert await backend(preview_page,'backend.isAdmin()&&backend.getProfile()===null')
                        else:
                            await select_russia(preview_page,server,'free')
                            assert await backend(preview_page,'!backend.isAdmin()&&backend.getProfile()===null')
                        preview_id=await finish_map(preview_page)
                        assert db('results/'+preview_id) is None
                        assert set(db('results') or {})==original_result_ids
                        assert await preview_page.locator('#leader_hard tr[data-attempt-id]').count()>0
                        record['excludedRoles'].append({'role':role,'noStudentProfile':True,'canReadLeaderboard':True,'inputPieces':89,'publicWrites':0})
                    except Exception:
                        await preview_page.screenshot(path=str(output/f'failure-{browser_name}-{role}.png'),full_page=True)
                        raise
                    finally:await preview.close()
                assert not record['pageErrors'],record['pageErrors']
                assert not record['unexpectedRequests'],record['unexpectedRequests']
                record['passed']=True
            except Exception as error:
                record['passed']=False;record['error']=str(error)
                if not page.is_closed():await page.screenshot(path=str(output/f'failure-{browser_name}.png'),full_page=True)
                raise
            finally:
                (output/f'emulator-{browser_name}.json').write_text(json.dumps(record,indent=2),encoding='utf8')
                await context.close();await browser.close()
    return record

async def first_offline_export(browser_name,output):
    """Load the locally cached library for the first time after transport loss."""
    group='XLSX offline '+browser_name
    keys=[]
    for index,difficulty in enumerate(('hard','medium','easy')):
        key='ui-offline-export-'+browser_name+'-'+difficulty;keys.append(key)
        db('results/'+key,{'fio':'=1+2 学生 '+difficulty,'group':group,'difficulty':difficulty,
            'time_ms':3599000,'elapsed_ms':7200123+index,'placed':89,'total':89,
            'timestamp':1767225600000+index,'user_agent':'isolated-offline-export',
            'participant_id':hashlib.sha256(key.encode()).hexdigest()},'PUT')
    result={'browser':browser_name,'scenario':'first-export-after-asset-network-loss','pageErrors':[],'libraryResponses':[],'libraryFailures':[]}
    with EmulatorSite() as server:
        async with async_playwright() as playwright:
            kwargs={'firefox_user_prefs':{'network.proxy.type':0}} if browser_name=='firefox' else {}
            browser=await getattr(playwright,browser_name).launch(headless=True,**kwargs)
            context=await browser.new_context(viewport={'width':1366,'height':1000},locale='ru-RU',service_workers='allow',accept_downloads=True)
            page=await context.new_page()
            page.on('pageerror',lambda error:result['pageErrors'].append(str(error)))
            page.on('response',lambda response:result['libraryResponses'].append({'status':response.status,'fromServiceWorker':response.from_service_worker}) if '/xlsx-0.20.3.full.min.js' in response.url else None)
            page.on('requestfailed',lambda request:result['libraryFailures'].append(request.failure) if '/xlsx-0.20.3.full.min.js' in request.url else None)
            try:
                await page.goto(server.base+'apps/puzzle.html?context=free',wait_until='domcontentloaded');await ready(page)
                await page.wait_for_function('Boolean(navigator.serviceWorker.controller)')
                await page.locator('#puzzleGroupFilter summary').click()
                await page.locator(f'#puzzleGroupFilterOptions input[value="{group}"]').check()
                for key in keys:await page.locator(f'tr[data-attempt-id="{key}"]').wait_for()
                assert await page.evaluate('typeof window.XLSX')=='undefined'
                vendor=server.base+'assets/vendor/xlsx/xlsx-0.20.3.full.min.js'
                assert await page.evaluate('async url=>Boolean(await caches.match(url))',vendor)
                server.offline=True
                # The independent HTTP client bypasses the worker, proving that
                # every same-origin request really receives the transport fault.
                request=urllib.request.Request(server.base+'uncached-offline-proof?nonce='+str(time.time_ns()))
                try:
                    with urllib.request.build_opener(urllib.request.ProxyHandler({})).open(request,timeout=10):
                        raise AssertionError('Site transport unexpectedly available')
                except urllib.error.HTTPError as error:assert error.code==503
                if browser_name!='webkit':await context.set_offline(True)
                result['transport']='browser-offline-and-site503' if browser_name!='webkit' else 'site-asset503'
                result['navigatorOnline']=await page.evaluate('navigator.onLine')
                if browser_name!='webkit':assert result['navigatorOnline'] is False
                result['independentSite503Proof']=True
                async with page.expect_download() as received:await page.locator('#puzzleLeaderboardExport').click()
                path=output/f'first-full-offline-{browser_name}.xlsx'
                await (await received.value).save_as(path)
                assert await page.evaluate('window.XLSX?.version')=='0.20.3'
                book=openpyxl.load_workbook(path)
                assert book.sheetnames==['Высокая','Средняя','Низкая']
                for index,sheet in enumerate(book.worksheets):
                    assert sheet.max_row==2 and sheet['B2'].data_type=='s'
                    assert sheet['B2'].value.startswith('=1+2 学生 ')
                    assert abs(sheet['D2'].value.total_seconds()*1000-(7200123+index))<1
                book.close()
                assert not result['pageErrors'],result['pageErrors']
                result.update({'passed':True,'firstLibraryExecutionOffline':True,'independentParser':'openpyxl','sheets':3,'rows':3})
            except Exception as error:
                result.update({'passed':False,'error':str(error)})
                result['diagnostic']=await page.evaluate("async()=>({online:navigator.onLine,library:typeof window.XLSX,worker:navigator.serviceWorker.controller?.scriptURL,caches:await caches.keys(),status:document.querySelector('#puzzleLeaderboardStatus').textContent})")
                await page.screenshot(path=str(output/f'first-full-offline-{browser_name}-failure.png'),full_page=True)
                raise
            finally:
                (output/f'first-full-offline-{browser_name}.json').write_text(json.dumps(result,indent=2),encoding='utf8')
                await context.close();await browser.close()
    return result

async def main():
    parser=argparse.ArgumentParser();parser.add_argument('--browsers',default='chromium');parser.add_argument('--offline-export-only',action='store_true');parser.add_argument('--output',type=Path,default=Path(os.environ.get('QA_OUT',str(ROOT.parent/'artifacts/rudn-puzzle-qa/leaderboard-1.3.7'))));args=parser.parse_args()
    args.output.mkdir(parents=True,exist_ok=True)
    for browser in args.browsers.split(','):
        if not args.offline_export_only:print(json.dumps(await run(browser,args.output)))
        print(json.dumps(await first_offline_export(browser,args.output)))

if __name__=='__main__':asyncio.run(main())
