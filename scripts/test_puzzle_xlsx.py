"""Independently read the actual application-generated XLSX bytes with openpyxl."""
import argparse
import datetime
import json
from pathlib import Path
import zipfile
import openpyxl

parser=argparse.ArgumentParser()
parser.add_argument('--artifacts',required=True,type=Path)
args=parser.parse_args()
expected=json.loads((args.artifacts/'xlsx-expected.json').read_text('utf8'))
reports=[]
for locale,copy in expected['copies'].items():
    path=args.artifacts/f'leaderboard-{locale}.xlsx'
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        assert all(b'<f>' not in archive.read(name) for name in archive.namelist() if name.startswith('xl/worksheets/') and name.endswith('.xml'))
    book=openpyxl.load_workbook(path,data_only=False)
    assert book.sheetnames==[copy[level] for level in ('hard','medium','easy')]
    count=0
    for level in ('hard','medium','easy'):
        sheet=book[copy[level]]
        rows=[r for r in expected['rows'] if r['difficulty']==level]
        assert sheet.max_row==len(rows)+1
        assert [cell.value for cell in sheet[1]]==[copy[k] for k in ('rank','fullName','group','time','dateUTC')]
        assert sheet.auto_filter.ref==f'A1:E{len(rows)+1}'
        for index,row in enumerate(rows,2):
            assert sheet.cell(index,1).value==index-1
            for column,key in ((2,'fio'),(3,'group')):
                cell=sheet.cell(index,column)
                assert cell.data_type=='s' and cell.value==row[key]
            elapsed=sheet.cell(index,4)
            assert elapsed.number_format=='[m]:ss.000'
            assert isinstance(elapsed.value,datetime.timedelta)
            assert abs(elapsed.value.total_seconds()*1000-row['elapsed_ms'])<1
            date=sheet.cell(index,5)
            assert isinstance(date.value,datetime.datetime)
            stamp=datetime.datetime.fromtimestamp(row['timestamp']/1000,datetime.timezone.utc).replace(tzinfo=None)
            assert abs((date.value-stamp).total_seconds())<0.001
            count+=1
    reports.append({'locale':locale,'sheets':book.sheetnames,'rows':count,'passed':True})
    book.close()
print(json.dumps({'independentParser':'openpyxl','passed':True,'reports':reports},indent=2))
